"""One-off helper: FEATURES_AND_PROGRESS.md -> FEATURES_AND_PROGRESS.docx"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE


def strip_md_bold(s: str) -> str:
    return s.replace("**", "")


def main():
    root = Path(__file__).resolve().parent
    md_path = root / "FEATURES_AND_PROGRESS.md"
    out_path = root / "FEATURES_AND_PROGRESS.docx"

    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("# ") and not stripped.startswith("##"):
            doc.add_heading(strip_md_bold(stripped[2:].strip()), level=0)
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(strip_md_bold(stripped[3:].strip()), level=1)
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("|") and "---" not in stripped:
            rows_raw = []
            while i < len(lines):
                row_line = lines[i].strip()
                if not row_line.startswith("|"):
                    break
                if row_line.replace("|", "").replace("-", "").replace(" ", "") == "":
                    i += 1
                    continue
                cells = [strip_md_bold(c.strip()) for c in row_line.strip("|").split("|")]
                rows_raw.append(cells)
                i += 1

            if rows_raw:
                ncol = max(len(r) for r in rows_raw)
                nrow = len(rows_raw)
                table = doc.add_table(rows=nrow, cols=ncol)
                table.style = "Table Grid"
                for ri in range(nrow):
                    for ci in range(ncol):
                        cell_text = rows_raw[ri][ci] if ci < len(rows_raw[ri]) else ""
                        table.cell(ri, ci).text = cell_text
            continue

        if stripped:
            p = doc.add_paragraph(strip_md_bold(line))
            p.paragraph_format.space_after = Pt(6)

        i += 1

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
